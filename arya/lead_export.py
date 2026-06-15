import os
import io
import base64
import pickle
from datetime import datetime
from email.mime.multipart import MIMEMultipart
from email.mime.text import MIMEText
from email.mime.base import MIMEBase
from email import encoders
from googleapiclient.discovery import build
from google.auth.transport.requests import Request
from dotenv import load_dotenv

load_dotenv(dotenv_path=os.path.join(os.path.dirname(os.path.abspath(__file__)), '..', '.env'), override=True)


def get_gmail_service():
    creds = None
    if os.path.exists('token_gmail.pickle'):
        with open('token_gmail.pickle', 'rb') as f:
            creds = pickle.load(f)
    if not creds or not creds.valid:
        if creds and creds.expired and creds.refresh_token:
            creds.refresh(Request())
    return build('gmail', 'v1', credentials=creds)


def leads_to_excel(leads: list, sheet_name: str = "Leads") -> bytes:
    """Convert leads list to Excel bytes using openpyxl."""
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment
    except ImportError:
        raise ImportError("openpyxl not installed")

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet_name

    headers = ["Name", "Email", "Phone", "Status", "Last Contact", "Next Follow-up", "No-show Count", "Notes"]
    header_fill = PatternFill(start_color="1a1a2e", end_color="1a1a2e", fill_type="solid")
    header_font = Font(color="00FF88", bold=True)

    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")
        ws.column_dimensions[chr(64 + col)].width = 20

    for row_i, lead in enumerate(leads, 2):
        for col_i, val in enumerate(lead[:8], 1):
            ws.cell(row=row_i, column=col_i, value=val)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def send_leads_excel_email(leads: list, to_email: str, filter_label: str = "All") -> str:
    """Generate Excel from leads and email it."""
    try:
        excel_bytes = leads_to_excel(leads, sheet_name=filter_label)
        filename = f"ARYA_leads_{filter_label}_{datetime.now().strftime('%Y%m%d')}.xlsx"

        service = get_gmail_service()
        msg = MIMEMultipart()
        msg['To'] = to_email
        msg['Subject'] = f"ARYA — {filter_label} Leads Export ({len(leads)} leads)"

        body = f"Hi Sagar,\n\nAttached are your {filter_label} leads ({len(leads)} total) exported from ARYA CRM.\n\nDate: {datetime.now().strftime('%d %b %Y')}\n\nBest,\nARYA"
        msg.attach(MIMEText(body, 'plain'))

        part = MIMEBase('application', 'octet-stream')
        part.set_payload(excel_bytes)
        encoders.encode_base64(part)
        part.add_header('Content-Disposition', f'attachment; filename="{filename}"')
        msg.attach(part)

        raw = base64.urlsafe_b64encode(msg.as_bytes()).decode()
        service.users().messages().send(userId='me', body={'raw': raw}).execute()

        return f"✅ Excel file sent to *{to_email}*\n📊 {len(leads)} {filter_label} leads exported\n📎 File: {filename}"
    except Exception as e:
        return f"❌ Export failed: {str(e)}"
